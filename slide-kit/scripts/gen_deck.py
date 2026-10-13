#!/usr/bin/env python3
"""Generate a .pptx deck in the Modernization corporate style (Template3).

Usage:
    python3 gen_deck.py spec.json -o deck.pptx

All geometry/colors/fonts below were extracted from the reference template
(13.33x7.5in, Lato, slate #445469 text, navy constellation cover art).
"""

from __future__ import annotations

import argparse
import json
import math
import re
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Length, Pt

# ---------------------------------------------------------------- constants
# Every value below is measured from Template3.pptx (pt on a 960x540 canvas).

SLIDE_W, SLIDE_H = 960.0, 540.0          # pt (16:9, 12192000x6858000 EMU)
BODY_TOP, BODY_BOTTOM = 88.0, 500.0      # usable body area on content slides
BODY_X, BODY_W = 46.7, 872.3             # body column (aligns with title box)
GAP = 14.0                               # vertical gap between body blocks

ASSETS = Path(__file__).resolve().parent.parent / "assets"

INK = "445469"         # theme dk1/dk2 = all body text (scheme tx1/tx2)
SLATE = "46546C"       # theme accent1 - flow boxes, connectors
BLUE = "1986AF"        # theme accent2 - table header, infobox base
TEAL = "46BE9C"        # theme accent3
LIME = "9CB631"        # theme accent4 - highlighted flow boxes
ORANGE = "FB8635"      # theme accent5 - dashed annotation lines
RED_ACC = "D53E23"     # theme accent6
WHITE = "FFFFFF"
INFOBOX_BG = "C9EAF7"  # accent2 lumMod20/lumOff80 - slide-6 info boxes
ZEBRA_L = "C9EAF7"     # table zebra
ZEBRA_R = "EAF6FC"
NOTE_RED = "FF0000"    # red bold-italic footnotes
PAGENUM_GRAY = "7F7F7F"  # bg1 lumMod50
RULE_GRAY = "A6A6A6"   # bg1 lumMod65 - tick left of the page number
GRAY = "7F7F7F"
BORDER = "BFBFBF"
LINK = "0066B3"        # square 1 blue, used for link text
SQ_BLUE = "0066B3"     # 3 accent squares under the title
SQ_ORANGE = "F36F21"
SQ_GREEN = "22B24C"
CHIP_G = "EFF5DA"      # lime tint
CHIP_B = "D9EDF8"      # blue tint
CHIP_P = "FDE4D0"      # orange tint
LEGEND_BG = "EDEFF2"
TIMELINE_LINE = "8A93A0"
TIMELINE_DOT = "45546B"

CODE_TEXT = "E8ECF1"   # code on dark slate panel
CODE_KEY = "E06C75"
CODE_STR = "98C379"
CODE_COM = "97A3B3"
CODE_NUM = "D19A66"

F_MAIN = "Lato Regular"   # template body/title typeface
F_PAGENUM = "Lato"
F_CODE = "Menlo"

MARKUP_RE = re.compile(r"(\*\*.+?\*\*|\[\[[a-z]+\|.+?\]\])", re.S)
NAMED_COLORS = {
    "red": RED_ACC, "green": "21B14B", "orange": ORANGE, "teal": BLUE,
    "tealdark": "0E5F7E", "blue": LINK, "gray": GRAY, "white": WHITE,
    "ink": INK, "lime": LIME, "slate": SLATE,
}

warnings: list[str] = []

SPEC_DIR = Path.cwd()  # set in main(); relative asset paths resolve against it


def warn(msg: str) -> None:
    warnings.append(msg)
    print(f"[warn] {msg}", file=sys.stderr)


# ---------------------------------------------------------------- primitives

def _rgb(hexstr: str) -> RGBColor:
    return RGBColor.from_string(hexstr.upper().lstrip("#"))


def add_picture(slide, path: str, x: float, y: float, w: float, h: float):
    return slide.shapes.add_picture(str(path), Pt(x), Pt(y), Pt(w), Pt(h))


def add_rect(slide, x, y, w, h, fill=None, line=None, line_w=1.0, rounded=False,
             radius=0.5, dash=None):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    sp = slide.shapes.add_shape(shape_type, Pt(x), Pt(y), Pt(w), Pt(h))
    if rounded:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = _rgb(fill)
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = _rgb(line)
        sp.line.width = Pt(line_w)
        if dash:
            _line_dash(sp, dash)
    return sp


def _line_dash(shape, dash: str):
    ln = shape.line._get_or_add_ln()
    d = ln.find(qn("a:prstDash"))
    if d is None:
        d = ln.makeelement(qn("a:prstDash"), {"val": dash})
        ln.append(d)
    else:
        d.set("val", dash)


def _set_arrowhead(shape, end_type="triangle"):
    ln = shape.line._get_or_add_ln()
    te = ln.find(qn("a:tailEnd"))
    if te is None:
        te = ln.makeelement(qn("a:tailEnd"), {"type": end_type, "w": "med", "len": "med"})
        ln.append(te)


def _add_line(slide, x1, y1, x2, y2, color=SLATE, line_w=1.0, dash=None, arrow=False):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2))
    conn.shadow.inherit = False
    conn.line.color.rgb = _rgb(color)
    conn.line.width = Pt(line_w)
    if dash:
        _line_dash(conn, dash)
    if arrow:
        _set_arrowhead(conn)
    return conn


def parse_markup(text: str):
    """'**bold**' and '[[color|text]]' -> [(text, bold, color_or_None)]"""
    out, pos = [], 0
    for m in MARKUP_RE.finditer(text):
        if m.start() > pos:
            out.append((text[pos:m.start()], False, None))
        tok = m.group(0)
        if tok.startswith("**"):
            out.append((tok[2:-2], True, None))
        else:
            inner = tok[2:-2]
            name, _, val = inner.partition("|")
            out.append((val, False, NAMED_COLORS.get(name, None)))
        pos = m.end()
    if pos < len(text):
        out.append((text[pos:], False, None))
    return out or [(text, False, None)]


def add_text(slide, x, y, w, h, content, size=10, color=INK, bold=False, italic=False,
             font=F_MAIN, align="left", anchor="top", line_spacing=None,
             space_after=0, wrap=True, shrink=False, inset=False):
    """content: str (markup allowed) or list[str] paragraphs.

    inset=True keeps PowerPoint's default 7.2/3.6pt box insets (used by the
    template's placeholder-equivalent text); default is zero margins.
    """
    tb = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    if inset:
        tf.margin_left = tf.margin_right = Pt(7.2)
        tf.margin_top = tf.margin_bottom = Pt(3.6)
    else:
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {
        "top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM,
    }[anchor]
    paras = content if isinstance(content, list) else [content]
    for i, ptext in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {
            "left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
            "right": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY,
        }[align]
        if line_spacing:
            p.line_spacing = line_spacing
        if space_after:
            p.space_after = Pt(space_after)
        for txt, b, c in parse_markup(str(ptext)):
            r = p.add_run()
            r.text = txt
            r.font.size = Pt(size)
            r.font.name = font
            r.font.bold = bold or b
            r.font.italic = italic
            r.font.color.rgb = _rgb(c or color)
    return tb


def set_bullet(par, char: str, color=INK, font="Arial", indent_pt=22.5, marL_pt=None):
    """Real bullet char, template style: text at marL, bullet hanging indent_pt
    before it (template uses marL 22.5/58.5/94.5pt with a fixed 22.5pt hang)."""
    marL = marL_pt if marL_pt is not None else indent_pt
    pPr = par._p.get_or_add_pPr()
    pPr.set("marL", str(int(marL * 12700)))
    pPr.set("indent", str(-int(indent_pt * 12700)))
    buClr = pPr.makeelement(qn("a:buClr"), {})
    srgb = buClr.makeelement(qn("a:srgbClr"), {"val": color})
    buClr.append(srgb)
    buFont = pPr.makeelement(qn("a:buFont"), {"typeface": font})
    buChar = pPr.makeelement(qn("a:buChar"), {"char": char})
    for el in (buClr, buFont, buChar):
        pPr.append(el)


def est_lines(text: str, width_pt: float, size: float, char_w: float = 0.52) -> int:
    """Rough wrapped-line estimate for Lato-like text."""
    text = re.sub(r"\*\*|\[\[[a-z]+\||\]\]", "", str(text))
    cpl = max(8, width_pt / (size * char_w))
    total = 0
    for chunk in text.split("\n"):
        total += max(1, math.ceil(len(chunk) / cpl))
    return total


def text_w(text: str, size: float) -> float:
    clean = re.sub(r"\*\*|\[\[[a-z]+\||\]\]", "", str(text))
    return len(clean) * size * 0.5


# ---------------------------------------------------------------- chrome
# Content-slide furniture measured on layout2: eyebrow + 32pt title + 3 accent
# squares + page number with gray tick (bottom-right). No logo is ever drawn.

EYEBROW_POS = (30.7, 11.8, 388.1, 11.6)
TITLE_POS = (30.7, 25.8, 872.3, 33.5)
SQUARES = ((33.8, 63.3), (48.3, 63.3), (62.8, 63.3))   # 8.6x7.2 each
SQ_SIZE = (8.6, 7.2)
PAGENUM_POS = (931.6, 508.9, 28.6, 28.8)
RULE_POS = (930.4, 512.1, 0.7, 21.6)
NOTE_BOX = (8.9, 460.0, 942.0, 78.4)   # bottom-anchored, grows upward


def content_chrome(slide, prs, title: str, eyebrow: str, page_no: int | None):
    add_text(slide, *EYEBROW_POS[:2], EYEBROW_POS[2], EYEBROW_POS[3], eyebrow,
             size=12, color=INK, inset=True, wrap=False)
    add_text(slide, *TITLE_POS, title, size=32, color=INK, inset=True,
             line_spacing=0.9, wrap=False)
    for (sx, sy), col in zip(SQUARES, (SQ_BLUE, SQ_ORANGE, SQ_GREEN)):
        add_rect(slide, sx, sy, SQ_SIZE[0], SQ_SIZE[1], fill=col, rounded=True,
                 radius=0.167)
    if page_no is not None:
        add_rect(slide, *RULE_POS, fill=RULE_GRAY)
        add_text(slide, *PAGENUM_POS, str(page_no), size=10, color=PAGENUM_GRAY,
                 font=F_PAGENUM, align="center", anchor="middle")


# ---------------------------------------------------------------- block renderers

def render_heading(slide, x, y, w, h, block):
    add_text(slide, x, y, w, min(h, 26), block.get("text", ""), size=block.get("size", 18),
             color=block.get("color", INK), bold=True, anchor="middle")


def render_bullets(slide, x, y, w, h, block):
    """Template bullet text: 18pt lvl0 / 16pt lvl1 / 14pt lvl2, marL 22.5pt hanging."""
    base = block.get("size", 18)
    scale = base / 18.0
    items = block["items"]
    tb = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    first = True
    for it in items:
        if isinstance(it, str):
            it = {"t": it}
        lvl = int(it.get("lvl", 0))
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(it.get("space_after", 6 if lvl == 0 else 2))
        p.line_spacing = 1.0
        set_bullet(p, "\u2022", color=it.get("color", INK), indent_pt=22.5,
                   marL_pt=22.5 + 36 * lvl)
        base_bold = bool(it.get("bold", False))
        size = it.get("size", (18, 16, 14)[min(lvl, 2)] * scale)
        for txt, b, c in parse_markup(it.get("t", "")):
            r = p.add_run()
            r.text = txt
            r.font.size = Pt(size)
            r.font.name = F_MAIN
            r.font.bold = base_bold or b
            r.font.color.rgb = _rgb(c or it.get("color", INK))


def fit_table(block, col_w: list[float], h_avail: float | None):
    """Pick a font size that (roughly) fits; returns (size, est_height)."""
    header = block.get("header", [])
    rows = block.get("rows", [])
    fracs = block.get("cols") or [1.0 / len(col_w)] * len(col_w)
    tot = sum(fracs)
    cw = [col_w[j] * fracs[j] / tot for j in range(len(col_w))]
    size = block.get("size")
    if size is not None:
        est = 0.0
        if header:
            est += 26
        for r in rows:
            lines = max(est_lines(c, cw[j] - 10, size) for j, c in enumerate(r))
            est += min(200, lines * size * 1.25 + 8)
        return size, est
    for cand in (11, 10, 9, 8):
        est = 26.0 if header else 0.0
        for r in rows:
            lines = max(est_lines(c, cw[j] - 10, cand) for j, c in enumerate(r))
            est += min(200, lines * cand * 1.25 + 8)
        if h_avail is None or est <= h_avail + 0.5:
            return cand, est
    return 8, est


def table_content_height(block, w: float) -> float:
    ncols = len(block.get("header", [])) or (
        len(block["rows"][0]) if block.get("rows") else 0)
    if ncols == 0:
        return 0.0
    fracs = block.get("cols") or [1.0 / ncols] * ncols
    tot = sum(fracs)
    col_w = [w * f / tot for f in fracs]
    _, est = fit_table(block, col_w, None)
    return est


def render_table(slide, x, y, w, h, block):
    header = block.get("header", [])
    rows = block.get("rows", [])
    ncols = len(header) or (len(rows[0]) if rows else 0)
    if ncols == 0:
        return
    nrows = len(rows) + (1 if header else 0)

    fracs = block.get("cols") or [1.0 / ncols] * ncols
    tot = sum(fracs)
    col_w = [w * f / tot for f in fracs]
    aligns = block.get("align") or ["left"] * ncols
    size, _ = fit_table(block, col_w, h)

    gf = slide.shapes.add_table(nrows, ncols, Pt(x), Pt(y), Pt(w), Pt(h))
    tbl = gf.table
    tbl.first_row = False
    tbl.horz_banding = False
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Pt(cw)

    def fill_cell(cell, text, size, color, bold, align, fill_hex, anchor_top=True):
        cell.fill.solid()
        cell.fill.fore_color.rgb = _rgb(fill_hex)
        cell.margin_left = cell.margin_right = Pt(4)
        cell.margin_top = cell.margin_bottom = Pt(2)
        cell.vertical_anchor = MSO_ANCHOR.TOP if anchor_top else MSO_ANCHOR.MIDDLE
        tf = cell.text_frame
        tf.word_wrap = True
        lines = str(text).split("\n")
        for i, lntext in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
                           "right": PP_ALIGN.RIGHT}[align]
            p.line_spacing = 1.0
            stripped = lntext.lstrip()
            bullet = None
            if stripped.startswith(("- ", "* ")):
                bullet, stripped = "\u2022", stripped[2:]
            elif stripped.startswith("o "):
                bullet, stripped = "\u25cb", stripped[2:]
            if bullet:
                set_bullet(p, bullet, color=color, indent_pt=8)
            for txt, b, c in parse_markup(stripped):
                r = p.add_run()
                r.text = txt
                r.font.size = Pt(size)
                r.font.name = F_MAIN
                r.font.bold = bold or b
                r.font.color.rgb = _rgb(c or color)

    ri = 0
    if header:
        tbl.rows[0].height = Pt(24)
        for j, htxt in enumerate(header):
            fill_cell(tbl.cell(0, j), htxt, size, WHITE, True, "left", BLUE)
        ri = 1
    for i, row in enumerate(rows):
        zebra = ZEBRA_L if i % 2 == 0 else ZEBRA_R
        lines = max(est_lines(c, col_w[j] - 10, size) for j, c in enumerate(row))
        tbl.rows[ri + i].height = Pt(min(200, lines * size * 1.25 + 8))
        for j, ctxt in enumerate(row):
            if str(ctxt).strip() == "^":     # spanned cell – keep empty, merged later
                continue
            fill_cell(tbl.cell(ri + i, j), ctxt, size, INK, False,
                      aligns[j] if j < len(aligns) else "left", zebra)

    # vertical merge: '^' merges the cell above with all consecutive '^' cells below
    if any("^" in str(c) for row in rows for c in row):
        for j in range(ncols):
            start = None
            for i in range(len(rows) + 1):
                is_caret = (i < len(rows) and j < len(rows[i])
                            and str(rows[i][j]).strip() == "^")
                if is_caret and start is None:
                    start = i - 1
                elif not is_caret and start is not None:
                    r0 = start + (1 if header else 0)
                    r1 = (i - 1) + (1 if header else 0)
                    if r1 > r0:
                        try:
                            origin = tbl.cell(r0, j)
                            origin.merge(tbl.cell(r1, j))
                            origin.vertical_anchor = MSO_ANCHOR.MIDDLE
                        except Exception:
                            warn(f"merge failed at col {j} rows {r0}-{r1}")
                    start = None


def render_image(slide, x, y, w, h, block):
    path = Path(block["path"])
    if not path.is_absolute():
        path = SPEC_DIR / path
    if not path.is_file():
        warn(f"image not found: {path}")
        return
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(w / iw, h / ih)
    dw, dh = iw * scale, ih * scale
    ix, iy = x + (w - dw) / 2, y + (h - dh) / 2
    add_picture(slide, str(path), ix, iy, dw, dh)

    caption = block.get("caption")
    if caption:
        cy = min(iy + dh + 6, SLIDE_H - 22)
        add_text(slide, x, cy, w, 20, caption, size=14, color=INK, italic=True,
                 align="center", wrap=False)

    # Slide-11 style annotations: slate label + dashed orange leader line.
    for co in block.get("callouts", []):
        lx, ly = co["xy"]
        lw = co.get("w", text_w(co["text"], 18))
        lh = 26.0
        add_text(slide, lx, ly, lw + 6, lh, co["text"], size=co.get("size", 18),
                 color=INK, wrap=False)
        px, py = co["point"]
        if px > lx + lw:                     # label sits left of the target point
            ex, ey = lx + lw + 4, ly + lh / 2
        elif px < lx:                        # label sits right of the target point
            ex, ey = lx - 4, ly + lh / 2
        else:
            ex, ey = lx + lw / 2, ly - 4 if py < ly else ly + lh + 4
        _add_line(slide, ex, ey, px, py, color=co.get("color", ORANGE),
                  line_w=co.get("line_w", 1.5), dash="dash",
                  arrow=co.get("arrow", False))


CODE_TOKEN_RE = re.compile(
    r'"(?:[^"\\]|\\.)*"\s*:|"(?:[^"\\]|\\.)*"|\btrue\b|\bfalse\b|\bnull\b|'
    r"\b\d+(?:\.\d+)?\b|[A-Za-z_][\w.]*|\s+|.")


def _code_line_runs(line: str):
    runs = []
    comment = None
    ci = line.find("//")
    if ci >= 0:
        code_part, comment = line[:ci], line[ci:]
    else:
        code_part = line
    for m in CODE_TOKEN_RE.finditer(code_part):
        tok = m.group(0)
        if tok.startswith('"') and tok.rstrip().endswith(":"):
            runs.append((tok, CODE_KEY))
        elif tok.startswith('"'):
            runs.append((tok, CODE_STR))
        elif tok in ("true", "false", "null") or tok.replace(".", "").isdigit():
            runs.append((tok, CODE_NUM))
        else:
            runs.append((tok, CODE_TEXT))
    if comment:
        runs.append((comment, CODE_COM))
    return runs


def render_code(slide, x, y, w, h, block):
    """Dark slate rounded panels (theme accent1 background)."""
    panels = block.get("panels") or [block]
    n = len(panels)
    gap = 12.0
    pw = (w - gap * (n - 1)) / n
    size = block.get("size", 8)
    for k, panel in enumerate(panels):
        px = x + k * (pw + gap)
        card = add_rect(slide, px, y, pw, h, fill=SLATE, rounded=True, radius=0.055)
        card.text_frame.text = ""
        ty = y + 6
        if panel.get("title"):
            add_text(slide, px + 10, ty, pw - 20, 14, panel["title"], size=8.5,
                     color=CODE_TEXT, bold=True)
            ty += 16
        body = panel.get("code", "")
        tb = slide.shapes.add_textbox(Pt(px + 10), Pt(ty), Pt(pw - 20), Pt(y + h - ty - 6))
        tf = tb.text_frame
        tf.word_wrap = block.get("wrap", True)
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        lines = body.split("\n")
        line_h = size * 1.32
        needed = len(lines) * line_h
        if needed > y + h - ty - 6:
            warn(f"code panel '{panel.get('title', k)}' overflows: needs {needed:.0f}pt, has {y + h - ty - 6:.0f}pt")
        first = True
        for lntext in lines:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.line_spacing = Pt(line_h)
            p.space_after = Pt(0)
            if not lntext:
                r = p.add_run(); r.text = " "; r.font.size = Pt(size); r.font.name = F_CODE
                continue
            for txt, colr in _code_line_runs(lntext):
                r = p.add_run()
                r.text = txt
                r.font.size = Pt(size)
                r.font.name = F_CODE
                r.font.color.rgb = _rgb(colr)


def render_cards(slide, x, y, w, h, block):
    groups = block.get("groups") or []
    legend = block.get("legend")
    n = len(groups)
    if n == 0:
        return
    gap = 10.0
    cw = (w - gap * (n - 1)) / n

    # pre-compute chip layout per group so cards hug their content
    layouts = []
    for g in groups:
        title = g.get("title", "")
        tlines = est_lines(title, cw - 12, 8.5) if title else 0
        cy = 8 + (12 * tlines + 6 if title else 0)
        chips = []
        for chip in g.get("chips", []):
            if isinstance(chip, str):
                chip = {"t": chip}
            lvl = int(chip.get("lvl", 0))
            if chip.get("link"):
                ch = 13
            else:
                ch = max(16, est_lines(chip["t"], cw - 16 - lvl * 10 - 12, 8) * 10.2 + 7)
            chips.append((chip, ch, lvl))
            cy += ch + 5
        layouts.append((g, title, tlines, chips, cy))
    card_h = min(h - (34 if legend else 0), max(cy for *_, cy in layouts))

    for k, (g, title, tlines, chips, content_h) in enumerate(layouts):
        gx = x + k * (cw + gap)
        add_rect(slide, gx, y, cw, card_h, fill=WHITE, line=BORDER,
                 line_w=0.75, rounded=True, radius=0.05)
        cy = y + 8
        if title:
            add_text(slide, gx + 6, cy, cw - 12, 12 * tlines, title, size=8.5,
                     color=INK, bold=True, align="center")
            cy += 12 * tlines + 6
        for chip, ch, lvl in chips:
            ix = gx + 8 + lvl * 10
            iw = cw - 16 - lvl * 10
            if chip.get("link"):
                tb = add_text(slide, ix, cy, iw, ch, chip["t"], size=8,
                              color=LINK, bold=False, font=F_MAIN)
                for r in tb.text_frame.paragraphs[0].runs:
                    r.font.underline = True
            else:
                fill_hex = {"green": CHIP_G, "blue": CHIP_B, "peach": CHIP_P}.get(
                    chip.get("tone", "green"), CHIP_G)
                add_rect(slide, ix, cy, iw, ch, fill=fill_hex, rounded=True, radius=0.28)
                add_text(slide, ix + 6, cy, iw - 12, ch, chip["t"], size=8,
                         color=INK, anchor="middle")
            cy += ch + 5
        if cy > y + card_h + 1:
            warn(f"card group '{title[:24]}' overflows card by {cy - (y + card_h):.0f}pt")

    if legend:
        ly = y + card_h + 8
        if ly + 26 > y + h:
            ly = y + h - 26
        add_rect(slide, x, ly, w, 26, fill=LEGEND_BG, rounded=True, radius=0.16)
        add_text(slide, x + 10, ly, 120, 26, legend.get("label", "Color key"),
                 size=10, color=INK, bold=True, anchor="middle")
        lx = x + w * 0.32
        for sw in legend.get("swatches", []):
            hexc, label = sw
            add_rect(slide, lx, ly + 7, 42, 12, fill=hexc)
            add_text(slide, lx + 50, ly, 130, 26, label, size=10, color=INK,
                     anchor="middle")
            lx += 195


def render_infobox(slide, x, y, w, h, block):
    """Slide-6 style: light-blue boxes, 12pt text with bold lead-ins.

    spec: {"type":"infobox","cols":[[box,...],[box,...]]}
      box = {"title": "...", "text": "..."}  (renders "**title** text")
      or  {"paras": ["**lead** rest", ...]}
    """
    cols = block.get("cols") or ([block.get("boxes", [])] if block.get("boxes") else [])
    cols = [c for c in cols if c]
    if not cols:
        return
    n = len(cols)
    col_gap = block.get("col_gap", 51.5)
    row_gap = block.get("row_gap", 23.3)
    cw = (w - col_gap * (n - 1)) / n
    size = block.get("size", 12)
    line_h = size * 1.35
    pad = 8.0

    def box_paras(bx) -> list[str]:
        if bx.get("paras"):
            return bx["paras"]
        title, text = bx.get("title", ""), bx.get("text", "")
        if title and text:
            return [f"**{title}** {text}"]
        return [title or text or ""]

    def box_height(bx) -> float:
        lines = sum(est_lines(ptext, cw - 2 * pad, size, char_w=0.56)
                    for ptext in box_paras(bx))
        return lines * line_h + 2 * pad

    y_cur = y
    max_rows = max(len(c) for c in cols)
    for r in range(max_rows):
        row_h = 0.0
        for ci, col in enumerate(cols):
            if r >= len(col):
                continue
            bx = col[r]
            bh = float(bx["h"]) if bx.get("h") else box_height(bx)
            row_h = max(row_h, bh)
        for ci, col in enumerate(cols):
            if r >= len(col):
                continue
            bx = col[r]
            bh = float(bx["h"]) if bx.get("h") else box_height(bx)
            bx_x = x + ci * (cw + col_gap)
            add_rect(slide, bx_x, y_cur, cw, bh, fill=block.get("fill", INFOBOX_BG))
            paras = box_paras(bx)
            tb = add_text(slide, bx_x + pad, y_cur + pad, cw - 2 * pad, bh - 2 * pad,
                          paras, size=size, color=INK)
            for p in tb.text_frame.paragraphs:
                p.line_spacing = 1.0
                p.space_after = Pt(2)
        y_cur += row_h + row_gap
    if y_cur - row_gap > y + h + 1:
        warn(f"infobox stack overflows by {y_cur - row_gap - (y + h):.0f}pt")


def infobox_content_height(block, w: float) -> float:
    cols = block.get("cols") or ([block.get("boxes", [])] if block.get("boxes") else [])
    cols = [c for c in cols if c]
    if not cols:
        return 0.0
    n = len(cols)
    col_gap = block.get("col_gap", 51.5)
    row_gap = block.get("row_gap", 23.3)
    cw = (w - col_gap * (n - 1)) / n
    size = block.get("size", 12)
    line_h = size * 1.35
    pad = 8.0

    def box_paras(bx) -> list[str]:
        if bx.get("paras"):
            return bx["paras"]
        title, text = bx.get("title", ""), bx.get("text", "")
        if title and text:
            return [f"**{title}** {text}"]
        return [title or text or ""]

    def box_height(bx) -> float:
        lines = sum(est_lines(ptext, cw - 2 * pad, size, char_w=0.56)
                    for ptext in box_paras(bx))
        return lines * line_h + 2 * pad

    total = 0.0
    for r in range(max(len(c) for c in cols)):
        row_h = max((box_height(col[r]) for col in cols if r < len(col)), default=0.0)
        total += row_h + row_gap
    return total - row_gap


# ---------------------------------------------------------------- flow (slide-5 diagram)

FLOW_SIZE = (114.0, 61.7)   # template node box
FLOW_RADIUS = 0.167
FLOW_LINE_W = 2.2

SIDES = {"l": (0.0, 0.5), "r": (1.0, 0.5), "t": (0.5, 0.0), "b": (0.5, 1.0)}


def render_flow(slide, x, y, w, h, block):
    """Rounded flow boxes + arrowed connectors.

    nodes: {"id","x","y","w":114,"h":61.7,"tone":"outline|slate|lime","text","sub":[..],"size":16}
    edges: {"from","to","type":"straight|elbow","from_side","to_side","dash","color","arrow":true}
    Coordinates are absolute page pt (x/y measured on the 960x540 canvas).
    """
    nodes = {n["id"]: n for n in block.get("nodes", [])}

    def anchor(nid, side):
        n = nodes[nid]
        fx, fy = SIDES[side]
        return (n["x"] + n.get("w", FLOW_SIZE[0]) * fx,
                n["y"] + n.get("h", FLOW_SIZE[1]) * fy)

    # edges first so boxes paint on top of line ends
    for e in block.get("edges", []):
        a, b = e["from"], e["to"]
        na, nb = nodes[a], nodes[b]
        acx, acy = na["x"] + na.get("w", FLOW_SIZE[0]) / 2, na["y"] + na.get("h", FLOW_SIZE[1]) / 2
        bcx, bcy = nb["x"] + nb.get("w", FLOW_SIZE[0]) / 2, nb["y"] + nb.get("h", FLOW_SIZE[1]) / 2
        s_side = e.get("from_side") or ("r" if bcx >= acx else "l") if abs(bcx - acx) >= abs(bcy - acy) \
            else e.get("from_side") or ("b" if bcy >= acy else "t")
        t_side = e.get("to_side") or ("l" if bcx >= acx else "r") if abs(bcx - acx) >= abs(bcy - acy) \
            else e.get("to_side") or ("t" if bcy >= acy else "b")
        p1 = anchor(a, s_side)
        p2 = anchor(b, t_side)
        color = e.get("color", SLATE)
        lw = e.get("line_w", FLOW_LINE_W)
        dash = e.get("dash")
        arrow = e.get("arrow", True)
        if e.get("points"):
            pts = [p1] + [tuple(q) for q in e["points"]] + [p2]
        elif e.get("type", "elbow") == "straight":
            pts = [p1, p2]
        else:
            if s_side in ("l", "r"):            # leave horizontally
                midx = (p1[0] + p2[0]) / 2
                pts = [p1, (midx, p1[1]), (midx, p2[1]), p2]
            else:                               # leave vertically
                midy = (p1[1] + p2[1]) / 2
                pts = [p1, (p1[0], midy), (p2[0], midy), p2]
        for i in range(len(pts) - 1):
            (x1, y1), (x2, y2) = pts[i], pts[i + 1]
            if abs(x1 - x2) < 0.05 and abs(y1 - y2) < 0.05:
                continue
            seg = _add_line(slide, x1, y1, x2, y2, color=color, line_w=lw,
                            dash=dash, arrow=(arrow and i == len(pts) - 2))

    for n in block.get("nodes", []):
        nx, ny = n["x"], n["y"]
        nw, nh = n.get("w", FLOW_SIZE[0]), n.get("h", FLOW_SIZE[1])
        tone = n.get("tone", "slate")
        fill, ln_color, tcol = {
            "outline": (None, SLATE, INK),
            "slate": (SLATE, SLATE, WHITE),
            "lime": (LIME, LIME, WHITE),
        }.get(tone, (n.get("fill", SLATE), n.get("line", SLATE), n.get("text_color", WHITE)))
        box = add_rect(slide, nx, ny, nw, nh, fill=fill, line=ln_color,
                       line_w=FLOW_LINE_W, rounded=True, radius=FLOW_RADIUS)
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Pt(2)
        tf.margin_top = tf.margin_bottom = Pt(1)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        size = n.get("size", 16)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.line_spacing = 1.0
        for txt, b, c in parse_markup(n.get("text", "")):
            r = p.add_run()
            r.text = txt
            r.font.size = Pt(size)
            r.font.name = F_MAIN
            r.font.bold = b
            r.font.color.rgb = _rgb(c or tcol)
        for i, sub in enumerate(n.get("sub", [])):
            sp = tf.add_paragraph()
            sp.alignment = PP_ALIGN.CENTER
            sp.line_spacing = 1.0
            r = sp.add_run()
            r.text = sub
            r.font.size = Pt(n.get("sub_size", 11))
            r.font.name = F_MAIN
            r.font.italic = True
            r.font.color.rgb = _rgb(tcol)


def flow_content_height(block, w: float) -> float:
    nodes = block.get("nodes", [])
    if not nodes:
        return 0.0
    top = min(n["y"] for n in nodes)
    bot = max(n["y"] + n.get("h", FLOW_SIZE[1]) for n in nodes)
    return bot - top


def render_timeline(slide, x, y, w, h, block):
    """Slide-3 style: horizontal line, dots, bold labels above (SmartArt look).

    Line spans x0=91 -> x1=858pt on the page (template measurement); labels are
    sized to their text so they never wrap.
    """
    stages = block.get("stages", [])
    n = len(stages)
    if n == 0:
        return
    x0 = block.get("x0", 91.0)
    x1 = block.get("x1", 858.0)
    ly = block.get("y", y + h / 2)
    label_size = block.get("label_size", 18)
    _add_line(slide, x0, ly, x1, ly, color=block.get("color", TIMELINE_LINE), line_w=1.5)
    pad = min(90.0, (x1 - x0) / (2 * n))
    dot_r = block.get("dot_r", 4.5)
    for i, st in enumerate(stages):
        cx = x0 + pad + (x1 - x0 - 2 * pad) * (i / (n - 1)) if n > 1 else (x0 + x1) / 2
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Pt(cx - dot_r), Pt(ly - dot_r),
                                     Pt(2 * dot_r), Pt(2 * dot_r))
        dot.shadow.inherit = False
        dot.fill.solid()
        dot.fill.fore_color.rgb = _rgb(block.get("dot_color", TIMELINE_DOT))
        dot.line.fill.background()
        label_w = block.get("label_w") or min(text_w(st, label_size) + 12, 280.0)
        add_text(slide, cx - label_w / 2, ly - 14 - label_size * 1.35, label_w,
                 label_size * 1.35, st, size=label_size, color=INK, bold=True,
                 align="center", anchor="bottom", wrap=False)


def timeline_content_height(block, w: float) -> float:
    return block.get("h", 90.0)


def render_columns(slide, x, y, w, h, block):
    """Side-by-side child blocks, e.g. bullets left + image right."""
    cols = block.get("cols") or []
    weights = block.get("weights") or [1.0] * len(cols)
    tot = sum(weights)
    gap = block.get("gap", 14.0)
    fracs = [wt / tot for wt in weights]
    cx = x
    for col, fr in zip(cols, fracs):
        cwid = (w - gap * (len(cols) - 1)) * fr
        kind = col.get("type", "bullets")
        fn = BLOCK_RENDERERS.get(kind)
        if fn is None:
            warn(f"unknown column block type '{kind}' skipped")
        else:
            ch = float(col["h"]) if col.get("h") else h
            cy = y
            if not col.get("h"):               # center content-sized blocks in column
                nat = content_height(col, cwid)
                if nat is not None and nat < h:
                    cy = y + (h - nat) / 2
            fn(slide, cx, cy, cwid, ch, col)
        cx += cwid + gap


BLOCK_RENDERERS = {
    "heading": render_heading,
    "bullets": render_bullets,
    "table": render_table,
    "image": render_image,
    "code": render_code,
    "cards": render_cards,
    "columns": render_columns,
    "infobox": render_infobox,
    "flow": render_flow,
    "timeline": render_timeline,
}


def bullets_content_height(block, w: float) -> float:
    base = block.get("size", 18)
    scale = base / 18.0
    total = 0.0
    for it in block.get("items", []):
        if isinstance(it, str):
            it = {"t": it}
        lvl = int(it.get("lvl", 0))
        s = it.get("size", (18, 16, 14)[min(lvl, 2)] * scale)
        total += est_lines(it.get("t", ""), w - 30 - 36 * lvl, s) * s * 1.32
        total += it.get("space_after", 6 if lvl == 0 else 2)
    return total


def code_content_height(block, w: float) -> float:
    panels = block.get("panels") or [block]
    size = block.get("size", 8)
    line_h = size * 1.32
    max_lines = max(len(p.get("code", "").split("\n")) for p in panels)
    return 12 + (16 if any(p.get("title") for p in panels) else 0) + max_lines * line_h + 4


def cards_content_height(block, w: float) -> float:
    groups = block.get("groups") or []
    n = max(1, len(groups))
    cw = (w - 10.0 * (n - 1)) / n
    best = 0.0
    for g in groups:
        title = g.get("title", "")
        cy = 8 + (12 * est_lines(title, cw - 12, 8.5) + 6 if title else 0)
        for chip in g.get("chips", []):
            if isinstance(chip, str):
                chip = {"t": chip}
            lvl = int(chip.get("lvl", 0))
            if chip.get("link"):
                ch = 13
            else:
                ch = max(16, est_lines(chip["t"], cw - 16 - lvl * 10 - 12, 8) * 10.2 + 7)
            cy += ch + 5
        best = max(best, cy)
    return best + (34 if block.get("legend") else 0)


def content_height(block, w: float) -> float | None:
    """Natural content height of a block, or None when it should fill the space."""
    t = block.get("type", "bullets")
    if t == "heading":
        return 26.0
    if t == "bullets":
        return bullets_content_height(block, w)
    if t == "table":
        return table_content_height(block, w)
    if t == "code":
        return code_content_height(block, w)
    if t == "cards":
        return cards_content_height(block, w)
    if t == "infobox":
        return infobox_content_height(block, w)
    if t == "flow":
        return flow_content_height(block, w)
    if t == "timeline":
        return timeline_content_height(block, w)
    return None  # image / columns / unknown → fill allocated space


def render_body(slide, body, w_body=BODY_W, x0=BODY_X):
    blocks = body if isinstance(body, list) else [body]
    blocks = [b for b in blocks if b]
    if not blocks:
        return
    avail = BODY_BOTTOM - BODY_TOP
    gaps = GAP * (len(blocks) - 1)

    heights: list[float | None] = []
    known_total = 0.0
    for b in blocks:
        if b.get("h"):
            hh = float(b["h"])
            heights.append(hh)
            known_total += hh
            continue
        ch = content_height(b, w_body)
        if ch is not None and ch < avail * 0.9:
            heights.append(ch)
            known_total += ch
        else:
            heights.append(None)

    unknown = [i for i, hh in enumerate(heights) if hh is None]
    if unknown:
        share = max(60.0, (avail - gaps - known_total) / len(unknown))
        for i in unknown:
            heights[i] = share
    total = sum(hh for hh in heights) + gaps

    y = BODY_TOP
    if total < avail:                      # center an under-filled stack vertically
        y += (avail - total) / 2
    for b, bh in zip(blocks, heights):
        kind = b.get("type", "bullets")
        fn = BLOCK_RENDERERS.get(kind)
        if fn is None:
            warn(f"unknown block type '{kind}' skipped")
        else:
            fn(slide, x0, y, w_body, bh, b)
        y += bh + GAP


# ---------------------------------------------------------------- slide layouts
# Geometry from template slides 1/4/8/10 (cover & section) and 15 (thanks).

COVER_TITLE = (32.3, 182.9, 626.6, 117.9)
COVER_SUBTITLE = (33.6, 311.7, 749.9, 28.4)
THANKS_TITLE_Y, THANKS_TITLE_H = 228.7, 71.0


def _notes(slide, spec):
    if spec.get("notes"):
        slide.notes_slide.notes_text_frame.text = spec["notes"]


def slide_cover(prs, spec):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_picture(slide, ASSETS / "bg_cover.jpg", 0, 0, SLIDE_W, SLIDE_H)
    size = spec.get("size", 50)
    ty = spec.get("title_y", COVER_TITLE[1])
    add_text(slide, COVER_TITLE[0], ty, COVER_TITLE[2], COVER_TITLE[3],
             spec.get("title", "Deck title"), size=size, color=WHITE, bold=True,
             anchor="middle", inset=True, line_spacing=0.9)
    sub = spec.get("subtitle")
    if sub:
        add_text(slide, COVER_SUBTITLE[0], spec.get("subtitle_y", COVER_SUBTITLE[1]),
                 COVER_SUBTITLE[2], COVER_SUBTITLE[3], sub, size=14, color=WHITE,
                 inset=True)
    _notes(slide, spec)
    return slide


def slide_content(prs, spec, page_no, eyebrow_default=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, fill=WHITE)
    content_chrome(slide, prs, spec.get("title", ""),
                   spec.get("eyebrow", eyebrow_default), page_no)
    body = spec.get("body")
    if body:
        render_body(slide, body)
    note = spec.get("note")
    if note:  # red bold-italic footnote bottom-left (template slides 2/5/9/11/14)
        paras = note if isinstance(note, list) else note.split("\n")
        tb = add_text(slide, *NOTE_BOX, paras, size=spec.get("note_size", 14),
                      color=spec.get("note_color", NOTE_RED), italic=True, bold=True,
                      anchor="bottom", inset=True)
        for i, p in enumerate(tb.text_frame.paragraphs):
            if i > 0:
                set_bullet(p, "\u2022", color=spec.get("note_color", NOTE_RED),
                           indent_pt=14)
                for r in p.runs:
                    r.font.italic = True
                    r.font.bold = True
    _notes(slide, spec)
    return slide


def slide_thanks(prs, spec):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_picture(slide, ASSETS / "bg_thanks.jpg", 0, 0, SLIDE_W, SLIDE_H)
    add_text(slide, 0, spec.get("title_y", THANKS_TITLE_Y), SLIDE_W, THANKS_TITLE_H,
             spec.get("title", "THANK YOU!"), size=spec.get("size", 50), color=WHITE,
             bold=True, align="center", anchor="middle", line_spacing=0.9)
    sub = spec.get("subtitle")
    if sub:
        add_text(slide, 0, spec.get("title_y", THANKS_TITLE_Y) + THANKS_TITLE_H + 8,
                 SLIDE_W, 26, sub, size=14, color=WHITE, align="center")
    _notes(slide, spec)
    return slide


# ---------------------------------------------------------------- main

def build(spec: dict, out_path: Path):
    prs = Presentation()
    prs.slide_width = Pt(SLIDE_W)
    prs.slide_height = Pt(SLIDE_H)
    eyebrow_default = spec.get("eyebrow", "")

    for i, s in enumerate(spec.get("slides", [])):
        layout = s.get("layout", "content")
        if layout == "cover":
            slide_cover(prs, s)
        elif layout == "section":
            slide_cover(prs, s)            # same geometry, no page number
        elif layout == "content":
            slide_content(prs, s, i + 1, eyebrow_default)
        elif layout == "thanks":
            slide_thanks(prs, s)
        else:
            warn(f"slide {i + 1}: unknown layout '{layout}' skipped")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    print(f"saved {out_path}")


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("spec", type=Path, help="deck spec JSON")
    ap.add_argument("-o", "--output", type=Path, help="output .pptx path")
    args = ap.parse_args()

    spec = json.loads(args.spec.read_text(encoding="utf-8"))
    global SPEC_DIR
    SPEC_DIR = args.spec.resolve().parent
    out = args.output or Path(spec.get("output") or args.spec.with_suffix(".pptx").name)
    build(spec, out)
    if warnings:
        print(f"completed with {len(warnings)} warning(s)", file=sys.stderr)


if __name__ == "__main__":
    main()
